# -*- coding: utf-8 -*-
"""
生成《CodePilot 项目详解》Word 文档（可打印）。

内容以「文档块」形式集中定义在 doc_content.py 的 BLOCKS 里，
本文件只负责排版：封面、目录域、多级标题、代码块、表格、提示框。

用法: python gen_project_doc.py
输出: CodePilot项目详解.docx
"""

import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from doc_content import BLOCKS, DOC_SUBTITLE, DOC_TITLE, DOC_VERSION

CN_FONT = "微软雅黑"
MONO_FONT = "Consolas"
EN_FONT = "Calibri"
BODY = Pt(10.5)
ACCENT = RGBColor(0x1F, 0x3B, 0x63)
SUBTLE = RGBColor(0x5A, 0x5A, 0x5A)
CODE_COLOR = RGBColor(0x1A, 0x1A, 0x1A)
H1_COLOR = RGBColor(0x14, 0x2A, 0x47)


# ---------------------------------------------------------------- #
# 基础排版工具
# ---------------------------------------------------------------- #
def _set_run_font(run, size, bold=False, color=None, cn=CN_FONT, en=EN_FONT, italic=False):
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:eastAsia"), cn)


def para(doc, before=0, after=4, line=1.35, align=None, indent=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent is not None:
        pf.left_indent = Cm(indent)
    if align is not None:
        p.alignment = align
    return p


def shade(element, fill):
    """给段落或单元格加背景色。"""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def bottom_border(p, color="1F3B63", size=8):
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    pPr.append(borders)


def left_bar(p, color="1F3B63", size=18):
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), color)
    borders.append(left)
    pPr.append(borders)


# ---------------------------------------------------------------- #
# 富文本：支持 **加粗** 与 `代码`
# ---------------------------------------------------------------- #
def rich(p, text, size=BODY, color=None, cn=CN_FONT):
    tokens = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            r = p.add_run(token[2:-2])
            _set_run_font(r, size, bold=True, color=color or RGBColor(0x1A, 0x1A, 0x1A), cn=cn)
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            r = p.add_run(token[1:-1])
            _set_run_font(r, Pt(size.pt - 0.5), color=RGBColor(0xB0, 0x30, 0x60), cn=cn, en=MONO_FONT)
        else:
            r = p.add_run(token)
            _set_run_font(r, size, color=color or RGBColor(0x1A, 0x1A, 0x1A), cn=cn)


# ---------------------------------------------------------------- #
# 块渲染
# ---------------------------------------------------------------- #
def h1(doc, text, numbered=None):
    if numbered:
        doc.add_page_break()
    p = para(doc, before=14 if not numbered else 0, after=10)
    label = f"{numbered}  {text}" if numbered else text
    r = p.add_run(label)
    _set_run_font(r, Pt(19), bold=True, color=H1_COLOR)
    bottom_border(p, size=14)
    return p


def h2(doc, text):
    p = para(doc, before=13, after=6)
    r = p.add_run(text)
    _set_run_font(r, Pt(14.5), bold=True, color=ACCENT)
    return p


def h3(doc, text):
    p = para(doc, before=10, after=5)
    r = p.add_run(text)
    _set_run_font(r, Pt(12), bold=True, color=RGBColor(0x2B, 0x4C, 0x7A))
    return p


def body(doc, text):
    p = para(doc, after=5)
    rich(p, text)
    return p


def bullets(doc, items, numbered=False):
    for index, item in enumerate(items, 1):
        p = para(doc, after=3, line=1.3)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        marker = f"{index}. " if numbered else "▪ "
        r = p.add_run(marker)
        _set_run_font(r, BODY, bold=not numbered, color=ACCENT if not numbered else None)
        rich(p, item)


def code(doc, text, caption=None):
    if caption:
        p = para(doc, before=6, after=2)
        r = p.add_run(caption)
        _set_run_font(r, Pt(9.5), bold=True, color=SUBTLE)

    lines = text.strip("\n").split("\n")
    for line in lines:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.06
        pf.left_indent = Cm(0.4)
        _shd = OxmlElement("w:shd")
        _shd.set(qn("w:val"), "clear")
        _shd.set(qn("w:color"), "auto")
        _shd.set(qn("w:fill"), "F4F6F8")
        p._p.get_or_add_pPr().append(_shd)
        r = p.add_run(line if line.strip() else " ")
        _set_run_font(r, Pt(8.8), color=CODE_COLOR, cn=MONO_FONT, en=MONO_FONT)
    para(doc, after=6, line=1.0)


def table(doc, headers, rows, widths=None, caption=None):
    if caption:
        p = para(doc, before=6, after=3)
        r = p.add_run(caption)
        _set_run_font(r, Pt(9.5), bold=True, color=SUBTLE)

    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr = t.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        _set_run_font(r, Pt(9.5), bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(hdr[i]._tc.get_or_add_tcPr(), "1F3B63")

    for row_data in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row_data):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.18
            rich(p, str(text), size=Pt(9.5))

    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Cm(width)
    para(doc, after=7, line=1.0)
    return t


def note(doc, text, kind="tip"):
    """提示框。kind: tip / warn / danger / info"""
    styles = {
        "tip": ("EDF7ED", "1E6B33", "💡 要点"),
        "warn": ("FFF6E5", "8A5A00", "⚠ 注意"),
        "danger": ("FDECEC", "9B1C1C", "✖ 常见错误"),
        "info": ("EEF3FA", "1F3B63", "ℹ 说明"),
    }
    fill, color, label = styles[kind]
    p = para(doc, before=5, after=5, line=1.3)
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.right_indent = Cm(0.2)
    shade(p._p.get_or_add_pPr(), fill)
    left_bar(p, color)
    r = p.add_run(label + "  ")
    _set_run_font(r, Pt(9.5), bold=True, color=RGBColor.from_string(color))
    rich(p, text, size=Pt(9.8))


def qa(doc, question, answer_points, followups=None):
    """面试 Q&A 排版。"""
    p = para(doc, before=9, after=4)
    r = p.add_run("Q  " + question)
    _set_run_font(r, Pt(11.5), bold=True, color=RGBColor(0x14, 0x2A, 0x47))

    for item in answer_points:
        if isinstance(item, tuple) and item and item[0] == "code":
            code(doc, item[1], caption=item[2] if len(item) > 2 else None)
        elif isinstance(item, tuple) and item and item[0] == "note":
            note(doc, item[1], item[2] if len(item) > 2 else "tip")
        else:
            p = para(doc, after=3, line=1.32)
            p.paragraph_format.left_indent = Cm(0.5)
            rich(p, item)

    if followups:
        p = para(doc, before=3, after=3)
        r = p.add_run("可能追问：")
        _set_run_font(r, Pt(9.8), bold=True, color=SUBTLE)
        for item in followups:
            pp = para(doc, after=2, line=1.28)
            pp.paragraph_format.left_indent = Cm(0.9)
            pp.paragraph_format.first_line_indent = Cm(-0.4)
            rr = pp.add_run("→ ")
            _set_run_font(rr, Pt(9.5), bold=True, color=ACCENT)
            rich(pp, item, size=Pt(9.8))


# ---------------------------------------------------------------- #
# 封面 / 目录 / 页脚
# ---------------------------------------------------------------- #
def add_page_number_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("第 ")
    _set_run_font(r, Pt(8.5), color=SUBTLE)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "17")
    rpr.append(sz)
    run.append(rpr)
    fld.append(run)
    p._p.append(fld)
    r2 = p.add_run(" 页")
    _set_run_font(r2, Pt(8.5), color=SUBTLE)


def build():
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = EN_FONT
    st.font.size = BODY
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    add_page_number_footer(sec)

    # ---------- 封面 ----------
    para(doc, before=120, after=0)
    p = para(doc, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(DOC_TITLE)
    _set_run_font(r, Pt(30), bold=True, color=ACCENT)

    p = para(doc, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(DOC_SUBTITLE)
    _set_run_font(r, Pt(13), color=SUBTLE)

    p = para(doc, after=30, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(DOC_VERSION)
    _set_run_font(r, Pt(10.5), color=SUBTLE)

    for line, size in [
        ("架构设计 · 核心代码 · 数据库设计 · 缓存降级 · 面试准备", 12),
        ("", 8),
        ("基于 LangGraph 多 Agent 编排 + RAG 检索增强生成", 11),
        ("FastAPI · LangGraph · Chroma · PostgreSQL · Redis · React", 10.5),
    ]:
        p = para(doc, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(line)
        _set_run_font(r, Pt(size), color=SUBTLE)

    # ---------- 目录 ----------
    doc.add_page_break()
    p = para(doc, after=10)
    r = p.add_run("目  录")
    _set_run_font(r, Pt(18), bold=True, color=H1_COLOR)
    bottom_border(p, size=12)

    p = para(doc, after=6)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-2" \h \z \u')
    inner = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "右键选择“更新域”以生成目录（本页为自动目录域）"
    inner.append(t)
    fld.append(inner)
    p._p.append(fld)

    note(doc, "在 Word 中打开后，于目录处右键 → 更新域 → 更新整个目录，即可生成带页码的目录。", "info")

    # ---------- 正文 ----------
    for block in BLOCKS:
        kind = block[0]
        if kind == "h1":
            _, text, numbered = block
            h1(doc, text, numbered)
        elif kind == "h2":
            h2(doc, block[1])
        elif kind == "h3":
            h3(doc, block[1])
        elif kind == "p":
            body(doc, block[1])
        elif kind == "ul":
            bullets(doc, block[1], numbered=False)
        elif kind == "ol":
            bullets(doc, block[1], numbered=True)
        elif kind == "code":
            code(doc, block[1], caption=block[2] if len(block) > 2 else None)
        elif kind == "table":
            _, headers, rows, widths = block
            table(doc, headers, rows, widths)
        elif kind == "note":
            note(doc, block[1], block[2] if len(block) > 2 else "tip")
        elif kind == "qa":
            _, question, answers, followups = block
            qa(doc, question, answers, followups)
        elif kind == "pagebreak":
            doc.add_page_break()
        else:
            raise ValueError(f"未知块类型: {kind}")

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "CodePilot项目详解.docx")
    doc.save(out)
    print("已生成:", out)
    return out


if __name__ == "__main__":
    build()
